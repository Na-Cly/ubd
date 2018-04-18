from docx import Document
import re, sys, json
regex = re.compile(r'xx[A-Za-z0-9]{2}xx')
import jsonLoader
template = 'template.docx'

def compareText(text, repDict):
	for i in range(0, len(repDict)):
		if repDict[i]['code'] in text:
			text = text.replace(repDict[i]['code'],repDict[i]['answer'])
	return text
def replaceAll(document,repDict):
	for table in document.tables:
		for i in range(0,len(table.rows)-1):
			for j in range(0, len(table.columns)-1):
					storeText = str(table.cell(i,j).text)
					table.cell(i,j).text = compareText(storeText, repDict)
					
if len(sys.argv) >= 2:
	#json file
	json_file = open(sys.argv[1],'r')
	json_data = json.loads(json_file.read())
	repDict = jsonLoader.loadData(json_data)
	
for key in repDict.keys():
	document = Document(template)
	#print(repDict[key])
	replaceAll(document,repDict[key])
	document.save(str(key))
	print('Saved ' + str(key))